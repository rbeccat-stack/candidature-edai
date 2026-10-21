"""Ensembles scolaires privés sous contrat de l'académie de Lyon, chiffrés et priorisés pour Ed.ai.

Entrées : trois fichiers CSV de data.education.gouv.fr, téléchargés dans un même dossier.
  - fr-en-college-effectifs-niveau-sexe-lv.csv    (effectifs DEPP des collèges)
  - fr-en-lycee_gt-effectifs-niveau-sexe-lv.csv   (effectifs DEPP des lycées GT)
  - fr-en-annuaire-education.csv                  (coordonnées et sites)

Sorties :
  - un fichier Excel prêt à importer dans un CRM, une ligne par compte, avec les UAI ;
  - avec --html, les données de la page (const DATA et const S) mises à jour sur place.

Usage :
  python construire_liste.py DOSSIER_SOURCES [--rentree 2025] [--sortie fichier.xlsx] [--html ../index.html]

Dépendance : openpyxl (pip install openpyxl).
"""

import argparse
import csv
import json
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ACADEMIE = "10"  # code DEPP de l'académie de Lyon (Ain, Loire, Rhône)
TARIF_MATIERE = 6  # € par élève et par an pour une matière, tarif public
BUREAUX = (45.7487, 4.8436)  # bureaux d'Ed.ai, rue Chevreul, Lyon
BONUS_COLLEGE_ET_LYCEE = 1.25
BONUS_MOINS_20_KM = 1.3
BONUS_MOINS_50_KM = 1.15
TAILLE_TOP = 25

# Réseau Centre scolaire Notre-Dame (logo sur le site d'Ed.ai), relevé à la main.
# À vérifier sur csnd.fr : ces comptes sont gardés dans la liste mais sortis du top.
CLIENTS_PROBABLES = {
    ("NOTRE DAME DE MONGRE", "VILLEFRANCHE-SUR-SAONE"),
    ("NOTRE DAME DE BEL AIR", "TARARE"),
    ("NOTRE DAME", "VILLEFRANCHE-SUR-SAONE"),
    ("NOTRE DAME", "BELLEVILLE-EN-BEAUJOLAIS"),
    ("NOTRE DAME DU MAS", "VAL D'OINGT"),
    ("NOTRE DAME", "CLAVEISOLLES"),
}

FICHIER_COLLEGES = "fr-en-college-effectifs-niveau-sexe-lv.csv"
FICHIER_LYCEES = "fr-en-lycee_gt-effectifs-niveau-sexe-lv.csv"
FICHIER_ANNUAIRE = "fr-en-annuaire-education.csv"


def lire_csv(chemin):
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        yield from csv.DictReader(f, delimiter=";")


def derniere_rentree(dossier):
    return max(int(r["Rentrée scolaire"]) for r in lire_csv(dossier / FICHIER_COLLEGES))


def etablissements(dossier, rentree):
    """Une ligne par établissement et par niveau : un lycée polyvalent avec une section collège compte deux fois."""
    sources = [
        ("Collège", FICHIER_COLLEGES, "Code académie", "Code département", "nombre_eleves_total"),
        ("Lycée GT", FICHIER_LYCEES, "Code académie et assimilé", "Code département pays", "Nombre d'élèves"),
    ]
    lignes = []
    for niveau, fichier, col_academie, col_dep, col_eleves in sources:
        for r in lire_csv(dossier / fichier):
            if r["Rentrée scolaire"] != str(rentree) or r[col_academie] != ACADEMIE:
                continue
            if not r["Secteur"].startswith("PRIVE"):
                continue
            eleves = int(r[col_eleves] or 0)
            if eleves > 0:
                lignes.append({
                    "uai": r["UAI"], "patro": r["Patronyme"], "commune": r["Commune"],
                    "dep": r[col_dep][-2:], "niveau": niveau, "eleves": eleves,
                })
    return lignes


def annuaire(dossier, uais):
    """Première fiche de l'annuaire pour chaque UAI (le fichier contient des doublons)."""
    fiches = {}
    for r in lire_csv(dossier / FICHIER_ANNUAIRE):
        uai = r["Identifiant_de_l_etablissement"]
        if uai in uais and uai not in fiches:
            fiches[uai] = r
    return fiches


def distance_km(lat, lon):
    p1, p2 = math.radians(BUREAUX[0]), math.radians(lat)
    dp, dl = p2 - p1, math.radians(lon - BUREAUX[1])
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


def comptes(lignes, fiches):
    """Collège et lycée de même nom dans la même commune = un seul compte, un seul décideur."""
    groupes = {}
    for e in sorted(lignes, key=lambda e: (e["patro"], e["commune"])):
        groupes.setdefault((e["patro"], e["commune"]), []).append(e)
    resultat = []
    for (nom, commune), es in groupes.items():
        college = sum(e["eleves"] for e in es if e["niveau"] == "Collège")
        lycee = sum(e["eleves"] for e in es if e["niveau"] == "Lycée GT")
        niveaux = sorted({e["niveau"] for e in es})
        coords = [(float(fiches[e["uai"]]["latitude"]), float(fiches[e["uai"]]["longitude"]))
                  for e in es if e["uai"] in fiches and fiches[e["uai"]]["latitude"]]
        lat = round(sum(c[0] for c in coords) / len(coords), 5) if coords else None
        lon = round(sum(c[1] for c in coords) / len(coords), 5) if coords else None
        dist = distance_km(lat, lon) if coords else None
        web = next((fiches[e["uai"]]["Web"] for e in es if e["uai"] in fiches and fiches[e["uai"]]["Web"]), None)
        eleves = college + lycee
        arr = eleves * TARIF_MATIERE
        ensemble = len(niveaux) == 2
        score = arr * (BONUS_COLLEGE_ET_LYCEE if ensemble else 1)
        if dist is not None:
            score *= BONUS_MOINS_20_KM if dist < 20 else BONUS_MOINS_50_KM if dist < 50 else 1
        resultat.append({
            "nom": nom, "commune": commune, "dep": es[0]["dep"],
            "niveaux": " + ".join(niveaux), "eleves": eleves, "college": college, "lycee": lycee,
            "arr_1_matiere": arr, "client_probable": (nom, commune) in CLIENTS_PROBABLES,
            "ensemble": ensemble, "score": score, "lat": lat, "lon": lon, "web": web,
            "km_siege": round(dist, 1) if dist is not None else None, "dist": dist,
            "uai": ", ".join(e["uai"] for e in es),
        })
    # Classement : les clients probables restent dans la liste mais ne prennent pas de rang.
    resultat.sort(key=lambda c: -c["score"])
    candidats = [c for c in resultat if not c["client_probable"]]
    for rang, c in enumerate(candidats[:TAILLE_TOP], 1):
        c["rang"] = rang
    for c in resultat:
        c.setdefault("rang", 0)
    resultat.sort(key=lambda c: (c["rang"] == 0, c["rang"], -c["score"]))
    return resultat


def synthese(lignes, cs, rentree):
    top = [c for c in cs if c["rang"]]
    clients = [c for c in cs if c["client_probable"]]
    eleves = sum(c["eleves"] for c in cs)
    deps = {}
    for c in cs:
        deps[c["dep"]] = deps.get(c["dep"], 0) + 1
    return {
        "rentree": rentree,
        "perimetre": "Académie de Lyon (Ain, Loire, Rhône), privé sous contrat",
        "nb_etablissements_uai": len(lignes),
        "nb_comptes": len(cs),
        "comptes_par_dep": dict(sorted(deps.items())),
        "eleves_total": eleves,
        "part_college_pct": round(100 * sum(c["college"] for c in cs) / eleves),
        "arr_1_matiere_total": sum(c["arr_1_matiere"] for c in cs),
        "nb_college_et_lycee": sum(c["ensemble"] for c in cs),
        "top25_eleves": sum(c["eleves"] for c in top),
        "top25_arr_1_matiere": sum(c["arr_1_matiere"] for c in top),
        "top25_arr_3_matieres": 3 * sum(c["arr_1_matiere"] for c in top),
        "top25_a_moins_de_20km": sum(1 for c in top if c["dist"] is not None and c["dist"] < 20),
        "clients_probables": [c["nom"] for c in clients],
        "clients_probables_eleves": sum(c["eleves"] for c in clients),
    }


ENTETES = ["Rang", "Ensemble scolaire", "Commune", "Dép.", "Niveaux", "Élèves", "dont collège", "dont lycée GT",
           "€/an 1 matière", "€/an 3 matières (max)", "Km du siège Ed.ai", "Collège + lycée", "Client probable",
           "Score", "Site", "UAI (identifiant HubSpot)"]
LARGEURS = [6, 34, 26, 6, 18, 8, 12, 12, 14, 20, 16, 14, 14, 9, 50, 26]


def ligne_compte(c):
    return [c["rang"], c["nom"], c["commune"], c["dep"], c["niveaux"], c["eleves"], c["college"], c["lycee"],
            c["arr_1_matiere"], 3 * c["arr_1_matiere"], c["km_siege"], int(c["ensemble"]),
            int(c["client_probable"]), round(c["score"]), c["web"], c["uai"]]


def ecrire_excel(chemin, lignes, cs, fiches, s):
    fmt = lambda n: f"{n:,}".replace(",", " ")
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    for texte in [
        "Synthèse",
        f"Périmètre : {s['perimetre']}, rentrée {s['rentree']}",
        f"{s['nb_comptes']} ensembles scolaires ({s['nb_etablissements_uai']} collèges et lycées GT) · {fmt(s['eleves_total'])} élèves",
        f"Valeur annuelle à 1 matière ({TARIF_MATIERE} €) : {fmt(s['arr_1_matiere_total'])} €",
        f"Part des élèves au collège : {s['part_college_pct']} %",
        f"Top {TAILLE_TOP} : {fmt(s['top25_eleves'])} élèves · {fmt(s['top25_arr_1_matiere'])} €/an sur 1 matière",
        "Clients probables (réseau CSND, à vérifier) : " + ", ".join(s["clients_probables"]),
    ]:
        ws.append([texte])
    ws["A1"].font = Font(bold=True)

    for titre, sel in [("Top 25", [c for c in cs if c["rang"]]), ("Tous les comptes", cs),
                       ("Clients probables", [c for c in cs if c["client_probable"]])]:
        ws = wb.create_sheet(titre)
        ws.append(ENTETES)
        for c in sel:
            ws.append(ligne_compte(c))
        mise_en_forme(ws, LARGEURS)

    ws = wb.create_sheet("Détail par établissement")
    ws.append(["uai", "patro", "commune", "dep", "niveau", "eleves", "Type_contrat_prive", "Web"])
    for e in lignes:
        f = fiches.get(e["uai"], {})
        ws.append([e["uai"], e["patro"], e["commune"], e["dep"], e["niveau"], e["eleves"],
                   f.get("Type_contrat_prive"), f.get("Web") or None])
    mise_en_forme(ws, [11, 34, 26, 6, 10, 8, 40, 50])

    ws = wb.create_sheet("Méthode")
    for texte in [
        "Méthode",
        f"Source : effectifs DEPP collèges + lycées GT, rentrée {s['rentree']}, secteur privé sous contrat, académie de Lyon. "
        "Coordonnées et sites : annuaire de l'Éducation nationale.",
        "Compte = ensemble scolaire (même nom + même commune) : collège et lycée = un seul décideur.",
        f"Valeur annuelle = élèves × {TARIF_MATIERE} € (tarif public, 1 matière). 3 matières = borne haute avant dégressivité.",
        "Score = valeur annuelle × 1,25 si collège + lycée, × 1,3 à moins de 20 km des bureaux / × 1,15 à moins de 50 km.",
        "Client probable = établissements Notre-Dame des communes du réseau CSND (logo sur ed.ai). "
        "À vérifier sur csnd.fr. Exclus du top 25, gardés dans la liste.",
        "Non inclus : lycées professionnels.",
    ]:
        ws.append([texte])
    ws["A1"].font = Font(bold=True)
    wb.save(chemin)


def mise_en_forme(ws, largeurs):
    for i, l in enumerate(largeurs):
        ws.column_dimensions[chr(65 + i)].width = l
    for cellule in ws[1]:
        cellule.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def maj_page(chemin, cs, s):
    """Remplace les données de la page ; le reste du fichier ne bouge pas."""
    data = [{
        "nom": c["nom"], "commune": c["commune"], "dep": c["dep"], "niveaux": c["niveaux"],
        "eleves": c["eleves"], "arr_1_matiere": c["arr_1_matiere"], "client_probable": c["client_probable"],
        "top25": c["rang"] > 0, "rang": c["rang"], "lat": c["lat"], "lon": c["lon"], "web": c["web"] or "",
        "km_siege": c["km_siege"],
    } for c in cs]
    html = Path(chemin).read_text(encoding="utf-8")
    html = re.sub(r"const DATA = \[.*?\];\n", lambda _: f"const DATA = {json.dumps(data, ensure_ascii=False)};\n", html, count=1, flags=re.S)
    html = re.sub(r"const S = \{.*?\n\};\n", lambda _: f"const S = {json.dumps(s, ensure_ascii=False, indent=2)};\n", html, count=1, flags=re.S)
    # Valeurs affichées si le JavaScript ne tourne pas
    html = re.sub(r'(<span id="rentree">)\d+(</span>)', rf"\g<1>{s['rentree']}\g<2>", html, count=1)
    html = re.sub(r'(<em id="h-n">)\d+(</em>)', rf"\g<1>{s['nb_comptes']}\g<2>", html, count=1)
    Path(chemin).write_text(html, encoding="utf-8", newline="\n")


def main():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("sources", type=Path, help="dossier contenant les trois CSV de data.education.gouv.fr")
    p.add_argument("--rentree", type=int, help="rentrée scolaire (par défaut : la plus récente)")
    p.add_argument("--sortie", type=Path, default=Path("edai_prive_academie_lyon.xlsx"))
    p.add_argument("--html", type=Path, help="page à mettre à jour (index.html)")
    args = p.parse_args()

    rentree = args.rentree or derniere_rentree(args.sources)
    lignes = etablissements(args.sources, rentree)
    fiches = annuaire(args.sources, {e["uai"] for e in lignes})
    cs = comptes(lignes, fiches)
    s = synthese(lignes, cs, rentree)
    ecrire_excel(args.sortie, lignes, cs, fiches, s)
    print(f"Rentrée {rentree} : {len(lignes)} collèges et lycées GT, {s['nb_comptes']} comptes, "
          f"{s['eleves_total']} élèves -> {args.sortie}")
    if args.html:
        maj_page(args.html, cs, s)
        print(f"Données de la page mises à jour : {args.html}")


if __name__ == "__main__":
    main()
